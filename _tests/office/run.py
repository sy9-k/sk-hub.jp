"""SK Hub Office のテストを、まとめて動かす

  python _tests/office/run.py            ブラウザ（Edge か Chrome）を見えないまま開いて run.js を動かし、書き出したファイルを check.py で確かめる
  python _tests/office/run.py --only sheet   Doc・Sheet・Slide・drive（Google ドライブへの保存。テスト用のドライブを使う）のどれか 1 つだけ

書き出したファイルは _tests/office/out/ に残る（Word・Excel・PowerPoint・LibreOffice で開いて、見た目も確かめられる）。
NG があると終了コード 1。build も外の部品もいらない（lxml だけ。python-pptx があれば .pptx も開いて確かめる）
"""
import json
import os
import shutil
import subprocess
import sys
import tempfile
import threading
import time
from functools import partial
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, urlparse

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent.parent
OUT = Path(os.environ.get("OFFICE_TEST_OUT") or (HERE / "out"))
done = threading.Event()
report = {}

BROWSERS = [
    os.environ.get("OFFICE_TEST_BROWSER", ""),
    r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
    r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
    r"C:\Program Files\Google\Chrome\Application\chrome.exe",
    "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
    "/Applications/Microsoft Edge.app/Contents/MacOS/Microsoft Edge",
    shutil.which("google-chrome") or "", shutil.which("chromium") or "", shutil.which("msedge") or "",
]


class Handler(SimpleHTTPRequestHandler):
    def log_message(self, *args):
        pass

    def do_POST(self):
        u = urlparse(self.path)
        body = self.rfile.read(int(self.headers.get("Content-Length") or 0))
        if u.path == "/__save":
            name = Path(parse_qs(u.query).get("name", ["file"])[0]).name
            (OUT / name).write_bytes(body)
        elif u.path == "/__done":
            report.update(json.loads(body.decode("utf-8")))
            done.set()
        self.send_response(204)
        self.end_headers()


def make_third_party_xlsx():
    """ほかのソフト（openpyxl）が作った、グラフ入りの .xlsx を input/ に置く（run.js が読みこむ。公開されない）。openpyxl が無ければ作らない"""
    shutil.rmtree(HERE / "input", ignore_errors=True)
    try:
        from openpyxl import Workbook
        from openpyxl.chart import BarChart, LineChart, Reference
    except ImportError:
        return
    (HERE / "input").mkdir()
    wb = Workbook()
    ws = wb.active
    ws.title = "データ"
    for row in [("月", "売上", "費用"), ("1月", 120, 80), ("2月", 150, 90), ("3月", 170, 100)]:
        ws.append(row)
    bar = BarChart()
    bar.type = "col"
    bar.title = "月ごと"
    bar.add_data(Reference(ws, min_col=2, max_col=3, min_row=1, max_row=4), titles_from_data=True)
    bar.set_categories(Reference(ws, min_col=1, min_row=2, max_row=4))
    bar.width, bar.height = 15, 7.5
    ws.add_chart(bar, "E2")
    line = LineChart()
    line.title = "売上だけ"
    line.add_data(Reference(ws, min_col=2, min_row=1, max_row=4), titles_from_data=True)
    ws.add_chart(line, "E20")
    # 項目の列が離れている（C 列のデータに A 列の項目）グラフは、範囲にできないので読まない
    far = BarChart()
    far.add_data(Reference(ws, min_col=3, min_row=1, max_row=4), titles_from_data=True)
    far.set_categories(Reference(ws, min_col=1, min_row=2, max_row=4))
    ws.add_chart(far, "N2")
    wb.save(HERE / "input" / "third-party-chart.xlsx")


def main(args):
    global OUT
    sys.stdout.reconfigure(encoding="utf-8")
    only = args[args.index("--only") + 1] if "--only" in args else ""
    shutil.rmtree(OUT, ignore_errors=True)
    if OUT.exists():
        # Word などで書き出したファイルを開いたままだと消せない。そのときは別のフォルダに書き出す
        OUT = HERE / ("out-" + time.strftime("%Y%m%d-%H%M%S"))
        print(f"※ {HERE / 'out'} を消せないので、{OUT.name} に書き出します")
    OUT.mkdir()
    make_third_party_xlsx()
    server =ThreadingHTTPServer(("127.0.0.1", 0), partial(Handler, directory=str(ROOT)))
    threading.Thread(target=server.serve_forever, daemon=True).start()
    url = f"http://127.0.0.1:{server.server_address[1]}/_tests/office/?post=1" + (f"&only={only}" if only else "")
    exe = next((b for b in BROWSERS if b and Path(b).exists()), None)
    if not exe:
        print("Edge か Chrome が見つかりません（OFFICE_TEST_BROWSER にパスを入れてください）")
        return 2
    profile = tempfile.mkdtemp(prefix="sk-office-test-")
    browser = subprocess.Popen([exe, "--headless=new", "--disable-gpu", "--no-first-run", "--no-default-browser-check", f"--user-data-dir={profile}", "--window-size=1400,900", url],
                               stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    started = time.time()
    try:
        if not done.wait(300):
            print("時間内に終わりませんでした（5 分）")
            return 2
    finally:
        browser.terminate()
        try:
            browser.wait(10)
        except subprocess.TimeoutExpired:
            browser.kill()
        server.shutdown()
        shutil.rmtree(profile, ignore_errors=True)
        shutil.rmtree(HERE / "input", ignore_errors=True)

    results = report.get("results", [])
    ng = [r for r in results if not r["ok"]]
    print(f"■ ブラウザの中のテスト（{time.time() - started:.1f} 秒）")
    for r in ng:
        print(f"NG  {r['app']} / {r['fixture']} / {r['step']} — {r['msg']}")
    print(f"{len(results)} 件中 {len(ng)} 件が NG" if ng else f"すべて OK（{len(results)} 件）")
    if only == "drive":
        return 1 if ng else 0
    print("\n■ Python で確かめる（check.py）")
    sys.path.insert(0, str(HERE))
    import check
    code = check.main([str(OUT)])
    return 1 if ng or code else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
